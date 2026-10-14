# Copyright (C) 2026 C0 contributors
# SPDX-License-Identifier: GPL-3.0-or-later
import zipfile

from docx import Document
from docx.oxml import OxmlElement
import pytest

from czero_core.extraction import ExtractionError, extract_isolated

DOCX = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'


@pytest.mark.parametrize('tag', ['sdt', 'ins', 'fldSimple', 'customXml', 'txbxContent'])
def test_unhandled_visible_docx_text_is_quarantined_not_partially_accepted(tmp_path, tag):
    path = tmp_path / 'original.docx'
    document = Document(); paragraph = document.add_paragraph('Supported body evidence')
    wrapper = OxmlElement('w:' + tag)
    container = wrapper
    if tag == 'sdt':
        container = OxmlElement('w:sdtContent'); wrapper.append(container)
    run = OxmlElement('w:r'); text = OxmlElement('w:t'); text.text = 'Unsupported visible evidence'
    run.append(text); container.append(run); paragraph._p.append(wrapper); document.save(path)
    original = path.read_bytes()
    with pytest.raises(ExtractionError, match='unsupported_docx'):
        extract_isolated(path, DOCX)
    assert path.read_bytes() == original


def test_utf16_docx_entity_declaration_is_rejected_without_expansion(tmp_path):
    path = tmp_path / 'original.docx'
    xml = '<?xml version="1.0" encoding="UTF-16"?><!DOCTYPE document [<!ENTITY hidden "untrusted">]><document>&hidden;</document>'
    with zipfile.ZipFile(path, 'w') as archive:
        archive.writestr('word/document.xml', xml.encode('utf-16'))
    with pytest.raises(ExtractionError, match='unsupported_xml'):
        extract_isolated(path, DOCX)
    assert path.exists()


def test_unsupported_main_part_with_arbitrary_extension_is_rejected(tmp_path):
    path = tmp_path / 'original.docx'
    document = Document(); document.add_paragraph('Supported body')
    control = OxmlElement('w:sdt'); content = OxmlElement('w:sdtContent')
    paragraph = OxmlElement('w:p'); run = OxmlElement('w:r'); text = OxmlElement('w:t')
    text.text = 'Omitted evidence'; run.append(text); paragraph.append(run)
    content.append(paragraph); control.append(content); document.element.body.insert(1, control)
    ordinary = tmp_path / 'ordinary.docx'; document.save(ordinary)
    with zipfile.ZipFile(ordinary) as source, zipfile.ZipFile(path, 'w') as target:
        for item in source.infolist():
            data = source.read(item)
            if item.filename in ('[Content_Types].xml', '_rels/.rels'):
                data = data.replace(b'document.xml', b'document.bin')
            filename = 'word/document.bin' if item.filename == 'word/document.xml' else item.filename
            target.writestr(filename, data)
    assert 'Supported body' in Document(path).paragraphs[0].text
    with pytest.raises(ExtractionError, match='unsupported_docx'):
        extract_isolated(path, DOCX)
    assert path.exists()
