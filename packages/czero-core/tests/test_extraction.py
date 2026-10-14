# Copyright (C) 2026 C0 contributors
# SPDX-License-Identifier: GPL-3.0-or-later
import json
from pathlib import Path
from types import SimpleNamespace
import zipfile

import pytest
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject,NameObject,DecodedStreamObject
from docx import Document
from czero_core.extraction import (Extracted,ExtractionError,extract_file,
    extract_isolated,chunks,MAX_UPLOAD)

DOCX='application/vnd.openxmlformats-officedocument.wordprocessingml.document'


def pdf(path,texts,*,encrypted=False,image=False):
    writer=PdfWriter()
    for text in texts:
        page=writer.add_blank_page(width=612,height=792)
        font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
        resources=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):writer._add_object(font)})})
        if image:
            stream=DecodedStreamObject();stream.set_data(b'\x00');stream.update({NameObject('/Subtype'):NameObject('/Image')})
            image_reference=writer._add_object(stream)
            if image=='nested':
                form=DecodedStreamObject();form.set_data(b'/Im0 Do');form.update({NameObject('/Subtype'):NameObject('/Form'),NameObject('/Resources'):DictionaryObject({NameObject('/XObject'):DictionaryObject({NameObject('/Im0'):image_reference})})})
                image_reference=writer._add_object(form)
            resources[NameObject('/XObject')]=DictionaryObject({NameObject('/Im0'):image_reference})
        page[NameObject('/Resources')]=writer._add_object(resources)
        content=DecodedStreamObject();content.set_data(('BT /F1 12 Tf 72 720 Td ('+text+') Tj ET').encode('ascii'))
        page[NameObject('/Contents')]=writer._add_object(content)
    if encrypted:writer.encrypt('fixture-only')
    writer.write(path)


@pytest.mark.parametrize('kind',['text/plain','text/markdown'])
def test_utf8_text_roundtrip_offsets_and_deterministic_chunks(tmp_path,kind):
    path=tmp_path/'source';text='# First\n\ncafé 東京 evidence\n\n# Second\n'+('chunk evidence '*100)
    path.write_text(text,encoding='utf-8')
    result=extract_file(path,kind)
    assert result.text==text and extract_isolated(path,kind)==result
    segments=chunks(result,size=256,overlap=32)
    assert chunks(result,size=256,overlap=32)==segments
    assert segments[-1]['end']==len(text)
    for segment in segments:assert segment['text']==text[segment['start']:segment['end']]
    assert result.text_hash==extract_file(path,kind).text_hash


def test_pdf_pages_preserve_full_text_and_locations(tmp_path):
    path=tmp_path/'source.pdf';pdf(path,['Inventory evidence','Project evidence'])
    result=extract_isolated(path,'application/pdf')
    assert result.text=='Inventory evidence\n\nProject evidence'
    assert [result.text[l['start']:l['end']] for l in result.locations]==['Inventory evidence','Project evidence']
    assert [l['index'] for l in result.locations]==[1,2]


def test_docx_body_tables_headers_and_footer_text(tmp_path):
    path=tmp_path/'source.docx';doc=Document();doc.add_paragraph('Body café 東京')
    table=doc.add_table(rows=1,cols=2);table.cell(0,0).text='Asset';table.cell(0,1).text='Server'
    doc.sections[0].header.paragraphs[0].text='Header evidence';doc.sections[0].footer.paragraphs[0].text='Footer evidence';doc.save(path)
    result=extract_isolated(path,DOCX)
    assert all(v in result.text for v in ['Body café 東京','Asset','Server','Header evidence','Footer evidence'])
    assert {l['kind'] for l in result.locations}=={'paragraph','table-cell','header','footer'}
    assert all(result.text[l['start']:l['end']] for l in result.locations)


@pytest.mark.parametrize('payload,code',[(b'\xff','invalid_text'),(b'one\0two','invalid_text'),(b' \n','no_text')])
def test_failed_text_extraction_keeps_source_original(tmp_path,payload,code):
    path=tmp_path/'original';path.write_bytes(payload)
    with pytest.raises(ExtractionError,match=code):extract_isolated(path,'text/plain')
    assert path.read_bytes()==payload


@pytest.mark.parametrize('encrypted,image,code',[(True,False,'encrypted_pdf'),(False,True,'unsupported_scan'),(False,'nested','unsupported_scan')],ids=['encrypted','direct-scan','nested-scan'])
def test_unsupported_pdf_retains_original(tmp_path,encrypted,image,code):
    path=tmp_path/'original.pdf';pdf(path,[''],encrypted=encrypted,image=image);original=path.read_bytes()
    with pytest.raises(ExtractionError,match=code):extract_isolated(path,'application/pdf')
    assert path.read_bytes()==original


def test_invalid_pdf_error_does_not_echo_bytes_or_path(tmp_path):
    path=tmp_path/'private-name.pdf';path.write_bytes(b'private-document-value')
    with pytest.raises(ExtractionError) as e:extract_isolated(path,'application/pdf')
    assert str(e.value)=='invalid_pdf'
    assert path.read_bytes()==b'private-document-value'


@pytest.mark.parametrize('member,data,code',[('word/document.xml',b'x'*200000,'archive_limit'),('../escape.xml',b'x','invalid_docx'),('word/document.xml',b'<!DOCTYPE secret>','unsupported_xml'),('word/vbaProject.bin',b'x','invalid_docx'),('word/footnotes.xml',b'x','invalid_docx')],ids=['archive-bomb','traversal','dtd','macro','unsupported-notes'])
def test_docx_unsupported_and_bomb_parts_are_not_silently_lost(tmp_path,member,data,code):
    path=tmp_path/'original.docx'
    with zipfile.ZipFile(path,'w',compression=zipfile.ZIP_DEFLATED) as z:z.writestr(member,data)
    with pytest.raises(ExtractionError,match=code):extract_isolated(path,DOCX)
    assert path.exists()


def test_special_files_limits_unknown_format_and_chunk_bounds(tmp_path):
    path=tmp_path/'original';path.write_text('evidence');link=tmp_path/'link';link.symlink_to(path)
    with pytest.raises(ExtractionError,match='invalid_file'):extract_file(link,'text/plain')
    with pytest.raises(ExtractionError,match='unsupported_format'):extract_isolated(path,'application/x-executable')
    with path.open('wb') as f:f.truncate(MAX_UPLOAD+1)
    with pytest.raises(ExtractionError,match='too_large'):extract_file(path,'text/plain')
    for size,overlap in [(True,0),(255,0),(4097,0),(256,65),(256,-1)]:
        with pytest.raises(ValueError):chunks(Extracted('text',[]),size=size,overlap=overlap)


def test_isolated_parser_deadline_and_signals_have_safe_codes(tmp_path,monkeypatch):
    import subprocess
    import czero_core.extraction as module
    def timeout(*args,**kwargs):raise subprocess.TimeoutExpired('private-path',25)
    monkeypatch.setattr(module.subprocess,'run',timeout)
    with pytest.raises(ExtractionError,match='parser_timeout'):extract_isolated(tmp_path/'file','text/plain')
    monkeypatch.setattr(module.subprocess,'run',lambda *args,**kwargs:SimpleNamespace(returncode=-9,stdout=b'',stderr=b'private-path'))
    with pytest.raises(ExtractionError,match='resource_limit'):extract_isolated(tmp_path/'file','text/plain')


def test_parser_isolation_works_without_existing_bytecode_cache(tmp_path,monkeypatch):
    path=tmp_path/'source.pdf';pdf(path,['Cold parser evidence'])
    monkeypatch.setenv('PYTHONPYCACHEPREFIX',str(tmp_path/'empty-bytecode-cache'))
    monkeypatch.delenv('PYTHONDONTWRITEBYTECODE',raising=False)
    assert extract_isolated(path,'application/pdf').text=='Cold parser evidence'



def test_chunk_cap_rejects_overflow_without_partial_success():
    document=Extracted('x'*1_000_000,[])
    with pytest.raises(ExtractionError,match='too_large'):chunks(document,size=256,overlap=64)
